import io
import re
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import delete, func, or_, select
from sqlalchemy.orm import Session

from ..models import Contact, ContactInteraction, User, UserApplication
from ..schemas_applications import (
    ContactCreate,
    ContactImportConfirmPayload,
    ContactImportPreviewResult,
    ContactImportPreviewRow,
    ContactInteractionCreate,
    ContactInteractionRead,
    ContactRead,
    ContactUpdate,
    RelatedApplicationSummary,
)


def normalize_email(email: str | None) -> str | None:
    if not email:
        return None
    cleaned = email.strip().lower()
    return cleaned if cleaned else None


def normalize_linkedin(url: str | None) -> str | None:
    if not url:
        return None
    cleaned = url.strip().lower().rstrip("/")
    cleaned = re.sub(r"^https?://(www\.)?", "", cleaned)
    return cleaned if cleaned else None


def build_contact_read(contact: Contact) -> ContactRead:
    # Factual application counts and latest dates
    apps = contact.applications or []
    interactions = contact.interactions or []

    interview_statuses = {
        "interviewing",
        "1st_interview",
        "2nd_interview",
        "3rd_interview",
        "final_interview",
        "offer",
        "accepted",
    }
    response_statuses = interview_statuses | {"rejected", "withdrawn"}

    interviews_count = sum(1 for a in apps if a.status in interview_statuses)
    responses_count = sum(1 for a in apps if a.status in response_statuses and a.status != "no_answer")

    # Dates
    last_interaction_date = None
    last_contacted_at = None
    last_response_at = None
    has_responded = False

    for inter in interactions:
        idate = inter.interaction_date
        if not last_interaction_date or idate > last_interaction_date:
            last_interaction_date = idate

        is_response = (
            inter.type in {"email_received", "response_received", "interview"}
            or ("response" in (inter.summary or "").lower())
        )
        if is_response:
            if not last_response_at or idate > last_response_at:
                last_response_at = idate
            has_responded = True
        elif inter.type in {"email_sent", "linkedin_message", "phone_call", "meeting", "other"}:
            if not last_contacted_at or idate > last_contacted_at:
                last_contacted_at = idate

    if not has_responded and responses_count > 0:
        has_responded = True

    # Next action from linked apps (e.g. earliest upcoming follow-up)
    next_action = None
    upcoming_dates = [a.next_follow_up for a in apps if a.next_follow_up]
    if upcoming_dates:
        earliest_fu = min(upcoming_dates)
        next_action = f"Follow-up: {earliest_fu.strftime('%b %d, %Y')}"

    # Determine communication status: 'not_contacted' | 'contacted_no_response' | 'responded'
    if has_responded:
        communication_status = "responded"
    elif last_contacted_at:
        communication_status = "contacted_no_response"
    else:
        communication_status = "not_contacted"

    # Preferred method and email used from actual interaction history
    preferred_contact_method = None
    email_used = None
    for inter in interactions:
        if not preferred_contact_method:
            if "email" in inter.type:
                preferred_contact_method = "email"
            elif "linkedin" in inter.type:
                preferred_contact_method = "linkedin"
            elif "phone" in inter.type:
                preferred_contact_method = "phone"
        if inter.details and "@" in inter.details:
            m = re.search(r"[\w\.-]+@[\w\.-]+", inter.details)
            if m:
                email_used = m.group(0)

    if not email_used:
        email_used = contact.email

    related_apps = []
    for a in apps:
        title = a.custom_job_title or (a.job.title if a.job else f"Application #{a.id}")
        comp = a.custom_company_name or (a.job.company.name if a.job and a.job.company else "Company")
        related_apps.append(
            RelatedApplicationSummary(
                id=a.id,
                job_title=title,
                company_name=comp,
                status=a.status,
                applied_at=a.applied_at,
                source=a.source or "Compust",
            )
        )

    interactions_read = [
        ContactInteractionRead.model_validate(i) if not isinstance(i, ContactInteractionRead) else i
        for i in interactions
    ]

    return ContactRead(
        id=contact.id,
        user_id=contact.user_id,
        first_name=contact.first_name,
        last_name=contact.last_name,
        name=contact.name,
        email=contact.email,
        phone=contact.phone,
        company=contact.company,
        role=contact.role,
        linkedin_url=contact.linkedin_url,
        location=contact.location,
        notes=contact.notes,
        tags=contact.tags,
        created_at=contact.created_at,
        updated_at=contact.updated_at,
        applications_count=len(apps),
        responses_count=responses_count,
        interviews_count=interviews_count,
        last_interaction_date=last_interaction_date,
        last_contacted_at=last_contacted_at,
        last_response_at=last_response_at,
        has_responded=has_responded,
        next_action=next_action,
        communication_status=communication_status,
        preferred_contact_method=preferred_contact_method,
        email_used=email_used,
        interactions=interactions_read,
        related_applications=related_apps,
    )


def get_user_contacts(
    db: Session,
    user_id: int,
    search: str | None = None,
    company: str | None = None,
    tag: str | None = None,
) -> list[ContactRead]:
    query = select(Contact).where(Contact.user_id == user_id)

    if search:
        pattern = f"%{search.strip()}%"
        query = query.where(
            or_(
                Contact.name.ilike(pattern),
                Contact.first_name.ilike(pattern),
                Contact.last_name.ilike(pattern),
                Contact.email.ilike(pattern),
                Contact.company.ilike(pattern),
                Contact.role.ilike(pattern),
                Contact.notes.ilike(pattern),
            )
        )

    if company:
        query = query.where(Contact.company.ilike(f"%{company.strip()}%"))

    if tag:
        query = query.where(Contact.tags.ilike(f"%{tag.strip()}%"))

    query = query.order_by(Contact.updated_at.desc())
    contacts = list(db.scalars(query).all())
    return [build_contact_read(c) for c in contacts]


def get_contact_by_id(db: Session, user_id: int, contact_id: int) -> Contact | None:
    return db.scalar(
        select(Contact).where(
            Contact.id == contact_id,
            Contact.user_id == user_id,
        )
    )


def find_matching_contacts(
    db: Session,
    user_id: int,
    email: str | None = None,
    name: str | None = None,
    company: str | None = None,
    linkedin_url: str | None = None,
) -> list[dict]:
    """Finds contacts that might match given email, name, company or linkedin."""
    norm_email = normalize_email(email)
    norm_li = normalize_linkedin(linkedin_url)
    results = []

    user_contacts = list(
        db.scalars(select(Contact).where(Contact.user_id == user_id)).all()
    )

    for c in user_contacts:
        c_email = normalize_email(c.email)
        c_li = normalize_linkedin(c.linkedin_url)

        if norm_email and c_email and norm_email == c_email:
            results.append({
                "contact": build_contact_read(c),
                "match_reason": f"Exact email match ({c.email})",
                "confidence": "exact",
            })
            continue

        if norm_li and c_li and norm_li == c_li:
            results.append({
                "contact": build_contact_read(c),
                "match_reason": "Exact LinkedIn profile match",
                "confidence": "high",
            })
            continue

        if name and c.name and name.strip().lower() == c.name.strip().lower():
            if company and c.company and company.strip().lower() == c.company.strip().lower():
                results.append({
                    "contact": build_contact_read(c),
                    "match_reason": f"Name & Company match ({c.name} at {c.company})",
                    "confidence": "high",
                })
            else:
                results.append({
                    "contact": build_contact_read(c),
                    "match_reason": f"Name match ({c.name})",
                    "confidence": "medium",
                })

    return results


def create_contact(db: Session, user: User, data: ContactCreate) -> Contact:
    now = datetime.now(timezone.utc).replace(tzinfo=None)

    first_name = data.first_name.strip() if data.first_name else None
    last_name = data.last_name.strip() if data.last_name else None
    name = data.name.strip() if data.name else ""

    if not name and (first_name or last_name):
        name = " ".join(filter(None, [first_name, last_name]))

    contact = Contact(
        user_id=user.id,
        first_name=first_name,
        last_name=last_name,
        name=name or "Contact",
        email=data.email.strip() if data.email else None,
        phone=data.phone.strip() if data.phone else None,
        company=data.company.strip() if data.company else None,
        role=data.role.strip() if data.role else None,
        linkedin_url=data.linkedin_url.strip() if data.linkedin_url else None,
        location=data.location.strip() if data.location else None,
        notes=data.notes.strip() if data.notes else None,
        tags=data.tags.strip() if data.tags else None,
        created_at=now,
        updated_at=now,
    )
    db.add(contact)
    db.flush()

    if data.application_id:
        app = db.scalar(
            select(UserApplication).where(
                UserApplication.id == data.application_id,
                UserApplication.user_id == user.id,
            )
        )
        if app:
            app.contact_id = contact.id
            if not app.contact_name:
                app.contact_name = contact.name
            if not app.contact_email and contact.email:
                app.contact_email = contact.email

    # Optionally record initial interaction / communication history
    if data.contacted:
        method = (data.contact_method or "email").lower()
        if method == "email":
            itype = "email_sent"
            isummary = f"Initial email to {contact.name}"
            details_parts = []
            if data.email_used:
                details_parts.append(f"Email used: {data.email_used.strip()}")
            if data.notes:
                details_parts.append(data.notes.strip())
            idetails = "\n".join(details_parts) if details_parts else None
        elif method == "linkedin":
            itype = "linkedin_message"
            isummary = f"LinkedIn outreach to {contact.name}"
            idetails = data.notes.strip() if data.notes else None
        elif method == "phone":
            itype = "phone_call"
            isummary = f"Phone call with {contact.name}"
            idetails = data.notes.strip() if data.notes else None
        else:
            itype = "other"
            isummary = f"Initial contact with {contact.name}"
            idetails = data.notes.strip() if data.notes else None

        contact_inter = ContactInteraction(
            contact_id=contact.id,
            application_id=data.application_id,
            user_id=user.id,
            type=itype,
            interaction_date=data.contact_date or now,
            summary=isummary,
            details=idetails,
            created_at=now,
        )
        db.add(contact_inter)

        if data.response_received:
            resp_date = data.response_date or now
            resp_type = "email_received" if method == "email" else "meeting"
            resp_summary = f"Response received from {contact.name}"
            resp_details = data.response_notes.strip() if data.response_notes else None
            resp_inter = ContactInteraction(
                contact_id=contact.id,
                application_id=data.application_id,
                user_id=user.id,
                type=resp_type,
                interaction_date=resp_date,
                summary=resp_summary,
                details=resp_details,
                created_at=now,
            )
            db.add(resp_inter)

    db.commit()
    db.refresh(contact)
    return contact


def update_contact(
    db: Session,
    user: User,
    contact_id: int,
    data: ContactUpdate,
) -> Contact | None:
    contact = get_contact_by_id(db, user.id, contact_id)
    if not contact:
        return None

    now = datetime.now(timezone.utc).replace(tzinfo=None)

    if data.first_name is not None:
        contact.first_name = data.first_name.strip() if data.first_name else None
    if data.last_name is not None:
        contact.last_name = data.last_name.strip() if data.last_name else None
    if data.name is not None:
        contact.name = data.name.strip()
    if data.email is not None:
        contact.email = data.email.strip() if data.email else None
    if data.phone is not None:
        contact.phone = data.phone.strip() if data.phone else None
    if data.company is not None:
        contact.company = data.company.strip() if data.company else None
    if data.role is not None:
        contact.role = data.role.strip() if data.role else None
    if data.linkedin_url is not None:
        contact.linkedin_url = data.linkedin_url.strip() if data.linkedin_url else None
    if data.location is not None:
        contact.location = data.location.strip() if data.location else None
    if data.notes is not None:
        contact.notes = data.notes.strip() if data.notes else None
    if data.tags is not None:
        contact.tags = data.tags.strip() if data.tags else None

    # Record subsequent interaction or response on edit
    if data.contacted:
        method = (data.contact_method or "email").lower()
        itype = "email_sent" if method == "email" else ("linkedin_message" if method == "linkedin" else ("phone_call" if method == "phone" else "other"))
        contact_inter = ContactInteraction(
            contact_id=contact.id,
            application_id=None,
            user_id=user.id,
            type=itype,
            interaction_date=data.contact_date or now,
            summary=f"Contact interaction via {method.capitalize()}",
            details=f"Email: {data.email_used.strip()}" if data.email_used else None,
            created_at=now,
        )
        db.add(contact_inter)

    if data.response_received:
        resp_date = data.response_date or now
        resp_inter = ContactInteraction(
            contact_id=contact.id,
            application_id=None,
            user_id=user.id,
            type="email_received",
            interaction_date=resp_date,
            summary=f"Response received from {contact.name}",
            details=data.response_notes.strip() if data.response_notes else None,
            created_at=now,
        )
        db.add(resp_inter)

    contact.updated_at = now
    db.commit()
    db.refresh(contact)
    return contact


def delete_contact(db: Session, user: User, contact_id: int) -> bool:
    contact = get_contact_by_id(db, user.id, contact_id)
    if not contact:
        return False

    db.execute(delete(Contact).where(Contact.id == contact_id, Contact.user_id == user.id))
    db.commit()
    return True


def log_contact_interaction(
    db: Session,
    user: User,
    contact_id: int,
    data: ContactInteractionCreate,
) -> ContactInteraction | None:
    contact = get_contact_by_id(db, user.id, contact_id)
    if not contact:
        return None

    now = datetime.now(timezone.utc).replace(tzinfo=None)
    idate = data.interaction_date or now

    interaction = ContactInteraction(
        contact_id=contact_id,
        application_id=data.application_id,
        user_id=user.id,
        type=data.type or "email_sent",
        interaction_date=idate,
        summary=data.summary.strip(),
        details=data.details.strip() if data.details else None,
        created_at=now,
    )
    db.add(interaction)

    contact.updated_at = now
    db.commit()
    db.refresh(interaction)
    return interaction


def delete_contact_interaction(db: Session, user: User, interaction_id: int) -> bool:
    inter = db.scalar(
        select(ContactInteraction).where(
            ContactInteraction.id == interaction_id,
            ContactInteraction.user_id == user.id,
        )
    )
    if not inter:
        return False

    db.execute(delete(ContactInteraction).where(ContactInteraction.id == interaction_id))
    db.commit()
    return True


# ---------------------------------------------------------------------------
# XLSX Export & Smart Import for Contacts
# ---------------------------------------------------------------------------

def export_contacts_xlsx(db: Session, user_id: int) -> bytes:
    contacts = get_user_contacts(db, user_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacts"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_font = Font(name="Calibri", size=10)
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    headers = [
        "First Name",
        "Last Name",
        "Full Name",
        "Email",
        "Phone",
        "Company",
        "Role / Title",
        "LinkedIn URL",
        "Location",
        "Applications",
        "Responded",
        "Interviews",
        "Next Action",
        "Notes",
        "Tags",
    ]

    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, c in enumerate(contacts, start=2):
        row_values = [
            c.first_name or "",
            c.last_name or "",
            c.name,
            c.email or "",
            c.phone or "",
            c.company or "",
            c.role or "",
            c.linkedin_url or "",
            c.location or "",
            c.applications_count,
            "Yes" if c.has_responded else "No",
            c.interviews_count,
            c.next_action or "",
            c.notes or "",
            c.tags or "",
        ]
        ws.append(row_values)
        ws.row_dimensions[row_idx].height = 20

        is_even = row_idx % 2 == 0
        for col_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = row_font
            cell.border = thin_border
            if is_even:
                cell.fill = alt_fill

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(contacts) + 1}"

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


COLUMN_HEADER_SYNONYMS: dict[str, list[str]] = {
    "email": ["email", "email address", "e-mail", "courriel", "mail", "contact email", "e-mail-adresse", "kontakt email"],
    "name": ["name", "full name", "contact name", "nom", "nom complet", "person", "kontaktname", "vor- und nachname"],
    "first_name": ["first name", "firstname", "prenom", "prénom", "given name", "vorname"],
    "last_name": ["last name", "lastname", "surname", "family name", "nachname", "familienname"],
    "company": ["company", "company name", "organization", "organisation", "entreprise", "employer", "unternehmen", "firma", "arbeitgeber"],
    "role": ["role", "title", "job title", "position", "poste", "fonction", "stelle", "berufsbezeichnung"],
    "phone": ["phone", "phone number", "telephone", "téléphone", "mobile", "cell", "telefon", "telefonnummer", "handy", "rufnummer"],
    "linkedin_url": ["linkedin", "linkedin url", "profile", "linkedin profile", "profil linkedin", "linkedin-profil"],
    "location": ["location", "city", "country", "ville", "pays", "address", "ort", "standort", "stadt", "land"],
    "notes": ["notes", "note", "comments", "remarques", "description", "notizen", "bemerkungen", "anmerkungen"],
    "tags": ["tags", "tag", "labels", "categories", "mots-cles", "schlagwörter", "kategorien"],
}


def preview_contacts_xlsx(db: Session, user_id: int, file_bytes: bytes) -> ContactImportPreviewResult:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return ContactImportPreviewResult(
            total_detected=0,
            new_count=0,
            existing_count=0,
            duplicate_count=0,
            column_mapping={},
            preview_rows=[],
        )

    raw_headers = [str(c or "").strip().lower() for c in rows[0]]
    col_mapping: dict[str, int] = {}

    for col_idx, header in enumerate(raw_headers):
        for field, synonyms in COLUMN_HEADER_SYNONYMS.items():
            if field not in col_mapping and (header == field or header in synonyms):
                col_mapping[field] = col_idx

    # Fetch existing user contacts to identify duplicates
    existing_contacts = list(
        db.scalars(select(Contact).where(Contact.user_id == user_id)).all()
    )
    existing_by_email = {
        normalize_email(c.email): c for c in existing_contacts if c.email
    }
    existing_by_li = {
        normalize_linkedin(c.linkedin_url): c for c in existing_contacts if c.linkedin_url
    }
    existing_by_name_comp = {
        (c.name.strip().lower(), (c.company or "").strip().lower()): c
        for c in existing_contacts
        if c.name
    }

    preview_rows: list[ContactImportPreviewRow] = []
    seen_in_sheet: set[str] = set()

    for row in rows[1:]:
        if not any(row):
            continue

        def get_val(field: str) -> str | None:
            if field in col_mapping and col_mapping[field] < len(row):
                val = row[col_mapping[field]]
                return str(val).strip() if val is not None else None
            return None

        fn = get_val("first_name")
        ln = get_val("last_name")
        raw_name = get_val("name")
        if not raw_name and (fn or ln):
            raw_name = " ".join(filter(None, [fn, ln]))

        if not raw_name and not get_val("email"):
            continue

        name = raw_name or "Contact"
        email = get_val("email")
        phone = get_val("phone")
        company = get_val("company")
        role = get_val("role")
        li_url = get_val("linkedin_url")
        location = get_val("location")
        notes = get_val("notes")
        tags = get_val("tags")

        # Duplicate detection
        status: "new" | "existing" | "duplicate" = "new"
        match_id = None
        match_name = None

        norm_em = normalize_email(email)
        norm_li = normalize_linkedin(li_url)
        name_comp_key = (name.lower(), (company or "").lower())

        if norm_em and norm_em in existing_by_email:
            matched = existing_by_email[norm_em]
            status = "existing"
            match_id = matched.id
            match_name = matched.name
        elif norm_li and norm_li in existing_by_li:
            matched = existing_by_li[norm_li]
            status = "existing"
            match_id = matched.id
            match_name = matched.name
        elif name_comp_key in existing_by_name_comp and company:
            matched = existing_by_name_comp[name_comp_key]
            status = "duplicate"
            match_id = matched.id
            match_name = matched.name
        elif norm_em and norm_em in seen_in_sheet:
            status = "duplicate"

        if norm_em:
            seen_in_sheet.add(norm_em)

        preview_rows.append(
            ContactImportPreviewRow(
                first_name=fn,
                last_name=ln,
                name=name,
                email=email,
                phone=phone,
                company=company,
                role=role,
                linkedin_url=li_url,
                location=location,
                notes=notes,
                tags=tags,
                status=status,
                match_id=match_id,
                match_name=match_name,
            )
        )

    new_count = sum(1 for r in preview_rows if r.status == "new")
    existing_count = sum(1 for r in preview_rows if r.status == "existing")
    dup_count = sum(1 for r in preview_rows if r.status == "duplicate")

    return ContactImportPreviewResult(
        total_detected=len(preview_rows),
        new_count=new_count,
        existing_count=existing_count,
        duplicate_count=dup_count,
        column_mapping={f: raw_headers[idx] for f, idx in col_mapping.items() if idx < len(raw_headers)},
        preview_rows=preview_rows,
    )


def import_contacts_from_preview(
    db: Session,
    user: User,
    payload: ContactImportConfirmPayload,
) -> dict:
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    created_count = 0
    updated_count = 0
    skipped_count = 0

    mode = payload.mode  # skip_duplicates | update_existing | import_all

    for row in payload.rows:
        if row.status in {"existing", "duplicate"}:
            if mode == "skip_duplicates":
                skipped_count += 1
                continue
            elif mode == "update_existing" and row.match_id:
                contact = get_contact_by_id(db, user.id, row.match_id)
                if contact:
                    if row.phone and not contact.phone:
                        contact.phone = row.phone
                    if row.company and not contact.company:
                        contact.company = row.company
                    if row.role and not contact.role:
                        contact.role = row.role
                    if row.linkedin_url and not contact.linkedin_url:
                        contact.linkedin_url = row.linkedin_url
                    if row.location and not contact.location:
                        contact.location = row.location
                    if row.notes:
                        contact.notes = f"{contact.notes}\n{row.notes}" if contact.notes else row.notes
                    contact.updated_at = now
                    updated_count += 1
                    continue

        contact = Contact(
            user_id=user.id,
            first_name=row.first_name,
            last_name=row.last_name,
            name=row.name,
            email=row.email,
            phone=row.phone,
            company=row.company,
            role=row.role,
            linkedin_url=row.linkedin_url,
            location=row.location,
            notes=row.notes,
            tags=row.tags,
            created_at=now,
            updated_at=now,
        )
        db.add(contact)
        created_count += 1

    db.commit()
    return {
        "status": "success",
        "created_count": created_count,
        "updated_count": updated_count,
        "skipped_count": skipped_count,
        "total_processed": len(payload.rows),
    }
