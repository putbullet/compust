import logging
from pathlib import Path
from typing import Any
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..schemas_resume import resolve_section_title
from .resume_localization import format_skill_display, format_language_display

logger = logging.getLogger("compust.resume_export")


class ResumeExportError(Exception):
    pass


def hex_to_rgb(hex_str: str) -> tuple[int, int, int]:
    clean = hex_str.lstrip("#")
    if len(clean) == 3:
        clean = "".join(c * 2 for c in clean)
    if len(clean) != 6:
        return (37, 99, 235)  # default blue
    return (int(clean[0:2], 16), int(clean[2:4], 16), int(clean[4:6], 16))


def validate_pdf_file(pdf_path: Path) -> None:
    if not pdf_path.exists() or pdf_path.stat().st_size == 0:
        raise ResumeExportError("Generated PDF file is empty or missing.")
    try:
        reader = PdfReader(str(pdf_path))
    except Exception as exc:
        raise ResumeExportError(f"Generated PDF cannot be read: {exc}")
    if len(reader.pages) == 0:
        raise ResumeExportError("Generated PDF contains 0 pages.")
    total_text = "".join(page.extract_text() or "" for page in reader.pages)
    if len(total_text.strip()) < 20:
        raise ResumeExportError("Generated PDF contains insufficient extractable text.")


def render_template_pdf(
    structured_data: dict[str, Any],
    settings: dict[str, Any],
    dest_path: Path,
) -> None:
    dest_path.parent.mkdir(parents=True, exist_ok=True)

    # 1. Primary typesetting engine: RenderCV (Typst)
    try:
        from .rendercv_service import render_rendercv_artifacts
        render_rendercv_artifacts(structured_data, settings, dest_pdf_path=dest_path)
        validate_pdf_file(dest_path)
        return
    except Exception as exc:
        logger.warning(f"RenderCV rendering failed ({exc}). Falling back to ReportLab engine.")

    # 2. Resilient fallback: ReportLab engine
    doc_size = A4 if settings.get("document_size", "A4").upper() == "A4" else letter
    doc = SimpleDocTemplate(
        str(dest_path),
        pagesize=doc_size,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    theme_hex = settings.get("theme_color", "#2563eb")
    template = settings.get("template", "modern")
    font_size_base = float(settings.get("font_size", 10.5))

    primary_color = colors.HexColor(theme_hex)
    dark_slate = colors.HexColor("#0f172a")
    sub_slate = colors.HexColor("#475569")
    line_color = colors.HexColor("#cbd5e1")

    styles = getSampleStyleSheet()

    # Dynamic typography styles based on template
    is_classic = template == "classic"
    is_minimal = template == "minimal"
    is_technical = template == "technical"

    heading_font = "Times-Bold" if is_classic else "Helvetica-Bold"
    body_font = "Times-Roman" if is_classic else "Helvetica"
    oblique_font = "Times-Italic" if is_classic else "Helvetica-Oblique"

    name_style = ParagraphStyle(
        "NameStyle",
        parent=styles["Normal"],
        fontName=heading_font,
        fontSize=font_size_base + 9.5,
        leading=font_size_base + 13.5,
        textColor=dark_slate if not is_minimal else dark_slate,
        spaceAfter=2,
    )

    headline_style = ParagraphStyle(
        "HeadlineStyle",
        parent=styles["Normal"],
        fontName=body_font,
        fontSize=font_size_base + 0.5,
        leading=font_size_base + 4,
        textColor=primary_color if not is_minimal else sub_slate,
        spaceAfter=4,
    )

    contact_style = ParagraphStyle(
        "ContactStyle",
        parent=styles["Normal"],
        fontName=body_font,
        fontSize=font_size_base - 1.5,
        leading=font_size_base + 2,
        textColor=sub_slate,
        spaceAfter=6,
    )

    section_heading_style = ParagraphStyle(
        "SecHead",
        parent=styles["Normal"],
        fontName=heading_font,
        fontSize=font_size_base + 1.5,
        leading=font_size_base + 5,
        textColor=primary_color if not is_minimal else dark_slate,
        spaceBefore=8,
        spaceAfter=4,
    )

    item_title_style = ParagraphStyle(
        "ItemTitle",
        parent=styles["Normal"],
        fontName=heading_font,
        fontSize=font_size_base,
        leading=font_size_base + 3.5,
        textColor=dark_slate,
    )

    item_meta_style = ParagraphStyle(
        "ItemMeta",
        parent=styles["Normal"],
        fontName=oblique_font,
        fontSize=font_size_base - 1.0,
        leading=font_size_base + 2.5,
        textColor=sub_slate,
    )

    body_style = ParagraphStyle(
        "BodyStyle",
        parent=styles["Normal"],
        fontName=body_font,
        fontSize=font_size_base - 0.5,
        leading=font_size_base + 3,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=3,
    )

    bullet_style = ParagraphStyle(
        "BulletStyle",
        parent=styles["Normal"],
        fontName=body_font,
        fontSize=font_size_base - 0.5,
        leading=font_size_base + 3,
        textColor=colors.HexColor("#1e293b"),
        leftIndent=12,
        firstLineIndent=-8,
        spaceAfter=2,
    )

    story = []

    profile = structured_data.get("profile", {})
    visibility = settings.get("section_visibility", {})
    order = settings.get("section_order", [
        "summary", "experience", "education", "skills", "projects", "certifications", "languages", "custom_sections"
    ])

    # Header section
    if visibility.get("profile", True):
        full_name = profile.get("full_name") or "Curriculum Vitae"
        headline = profile.get("headline", "")
        contact_parts = []
        if profile.get("email"):
            contact_parts.append(profile["email"])
        if profile.get("phone"):
            contact_parts.append(profile["phone"])
        if profile.get("location"):
            contact_parts.append(profile["location"])
        if profile.get("linkedin"):
            contact_parts.append(profile["linkedin"])
        if profile.get("github"):
            contact_parts.append(profile["github"])
        if profile.get("website"):
            contact_parts.append(profile["website"])

        contact_line = " • ".join(contact_parts)

        story.append(Paragraph(full_name, name_style))
        if headline:
            story.append(Paragraph(headline, headline_style))
        if contact_line:
            story.append(Paragraph(contact_line, contact_style))

        if is_classic:
            story.append(HRFlowable(width="100%", thickness=1, color=dark_slate, spaceBefore=2, spaceAfter=6))
        elif is_minimal:
            story.append(HRFlowable(width="100%", thickness=0.5, color=line_color, spaceBefore=2, spaceAfter=6))
        else:
            story.append(HRFlowable(width="100%", thickness=1.5, color=primary_color, spaceBefore=2, spaceAfter=6))

    lang = (settings.get("language") or "en").lower()
    section_titles = settings.get("section_titles") or {}

    # Render sections according to section_order
    for sec_key in order:
        if not visibility.get(sec_key, True):
            continue

        resolved_heading = resolve_section_title(sec_key, section_titles.get(sec_key), lang=lang).upper()

        if sec_key == "summary" and profile.get("summary"):
            story.append(Paragraph(resolved_heading, section_heading_style))
            story.append(Paragraph(profile["summary"].replace("\n", "<br/>"), body_style))
            story.append(Spacer(1, 4))

        elif sec_key == "skills" and structured_data.get("skills"):
            skills = structured_data["skills"]
            if skills:
                story.append(Paragraph(resolved_heading, section_heading_style))
                # Group by category if available
                categories: dict[str, list[str]] = {}
                for s in skills:
                    cat = s.get("category", "General") or "General"
                    formatted = format_skill_display(s.get("name", ""), s.get("proficiency"), lang=lang)
                    if formatted:
                        categories.setdefault(cat, []).append(formatted)

                for cat, items in categories.items():
                    if len(categories) > 1 and cat != "General":
                        line = f"<b>{cat}:</b> " + ", ".join(items)
                    else:
                        line = ", ".join(items)
                    story.append(Paragraph(line, body_style))
                story.append(Spacer(1, 4))

        elif sec_key == "experience" and structured_data.get("experience"):
            experiences = structured_data["experience"]
            if experiences:
                story.append(Paragraph(resolved_heading, section_heading_style))
                for exp in experiences:
                    title = exp.get("title", "")
                    company = exp.get("company", "")
                    loc = exp.get("location", "")
                    start_d = exp.get("start_date", "")
                    end_d = exp.get("end_date", "")
                    date_range = f"{start_d} – {end_d}".strip(" –")

                    header_line = f"<b>{title}</b>" if title else ""
                    if company:
                        header_line += f" | {company}"
                    if loc:
                        header_line += f" ({loc})"

                    story.append(Paragraph(header_line, item_title_style))
                    if date_range:
                        story.append(Paragraph(date_range, item_meta_style))

                    desc = exp.get("description", "")
                    if desc:
                        for line in desc.splitlines():
                            line_clean = line.strip()
                            if not line_clean:
                                continue
                            if line_clean.startswith(("-", "•", "*")):
                                story.append(Paragraph(f"• {line_clean.lstrip('-•* ')}", bullet_style))
                            else:
                                story.append(Paragraph(line_clean, body_style))

                    highlights = exp.get("highlights", [])
                    for hl in highlights:
                        if hl and hl.strip():
                            story.append(Paragraph(f"• {hl.strip()}", bullet_style))

                    story.append(Spacer(1, 3))

        elif sec_key == "education" and structured_data.get("education"):
            educations = structured_data["education"]
            if educations:
                story.append(Paragraph(resolved_heading, section_heading_style))
                for edu in educations:
                    inst = edu.get("institution", "")
                    degree = edu.get("degree", "")
                    field = edu.get("field", "")
                    start_d = edu.get("start_date", "")
                    end_d = edu.get("end_date", "")
                    date_range = f"{start_d} – {end_d}".strip(" –")

                    edu_line = f"<b>{degree} in {field}</b>" if (degree and field) else f"<b>{degree or field or 'Degree'}</b>"
                    if inst:
                        edu_line += f" — {inst}"
                    story.append(Paragraph(edu_line, item_title_style))
                    if date_range:
                        story.append(Paragraph(date_range, item_meta_style))
                    if edu.get("description"):
                        story.append(Paragraph(edu["description"], body_style))
                    story.append(Spacer(1, 3))

        elif sec_key == "projects" and structured_data.get("projects"):
            projects = structured_data["projects"]
            if projects:
                story.append(Paragraph(resolved_heading, section_heading_style))
                for proj in projects:
                    name = proj.get("name", "")
                    techs = proj.get("technologies", "")
                    url = proj.get("url", "")
                    title_line = f"<b>{name}</b>"
                    if techs:
                        title_line += f" ({techs})"
                    story.append(Paragraph(title_line, item_title_style))
                    if url:
                        story.append(Paragraph(f"<u>{url}</u>", item_meta_style))
                    if proj.get("description"):
                        story.append(Paragraph(proj["description"], body_style))
                    story.append(Spacer(1, 3))

        elif sec_key == "certifications" and structured_data.get("certifications"):
            certs = structured_data["certifications"]
            if certs:
                story.append(Paragraph(resolved_heading, section_heading_style))
                for cert in certs:
                    name = cert.get("name", "")
                    issuer = cert.get("issuer", "")
                    date_str = cert.get("issue_date", "")
                    cert_line = f"• <b>{name}</b>"
                    if issuer:
                        cert_line += f" — {issuer}"
                    if date_str:
                        cert_line += f" ({date_str})"
                    story.append(Paragraph(cert_line, bullet_style))
                story.append(Spacer(1, 3))

        elif sec_key == "languages" and structured_data.get("languages"):
            langs = structured_data["languages"]
            if langs:
                story.append(Paragraph(resolved_heading, section_heading_style))
                lang_items = [
                    format_language_display(l.get("language", ""), l.get("proficiency"), lang=lang)
                    for l in langs
                    if l.get("language")
                ]
                lang_items = [li for li in lang_items if li]
                if lang_items:
                    story.append(Paragraph(" • ".join(lang_items), body_style))
                story.append(Spacer(1, 3))

        elif sec_key == "custom_sections" and structured_data.get("custom_sections"):
            custom_sections = structured_data["custom_sections"]
            for c_sec in custom_sections:
                title = (c_sec.get("title") or resolved_heading).strip().upper()
                items = c_sec.get("items", [])
                if title and items:
                    story.append(Paragraph(title, section_heading_style))
                    for it in items:
                        if it and it.strip():
                            story.append(Paragraph(f"• {it.strip()}", bullet_style))
                    story.append(Spacer(1, 3))
    story.append(Spacer(1, 3))

    doc.build(story)
    validate_pdf_file(dest_path)


def render_template_docx(
    structured_data: dict[str, Any],
    settings: dict[str, Any],
    dest_path: Path,
) -> None:
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # Set page margins
    for s in doc.sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)

    theme_hex = settings.get("theme_color", "#2563eb")
    r, g, b = hex_to_rgb(theme_hex)
    theme_rgb = RGBColor(r, g, b)

    profile = structured_data.get("profile", {})
    visibility = settings.get("section_visibility", {})
    order = settings.get("section_order", [
        "summary", "experience", "education", "skills", "projects", "certifications", "languages", "custom_sections"
    ])

    # Header
    if visibility.get("profile", True):
        name = profile.get("full_name") or "Curriculum Vitae"
        h = doc.add_heading(name, level=0)
        h.runs[0].font.size = Pt(22)
        h.runs[0].font.color.rgb = RGBColor(15, 23, 42)

        if profile.get("headline"):
            p_head = doc.add_paragraph()
            r_head = p_head.add_run(profile["headline"])
            r_head.font.size = Pt(12)
            r_head.font.color.rgb = theme_rgb
            r_head.bold = True

        contacts = [profile.get(k) for k in ["email", "phone", "location", "linkedin", "github", "website"] if profile.get(k)]
        if contacts:
            p_con = doc.add_paragraph()
            r_con = p_con.add_run(" | ".join(contacts))
            r_con.font.size = Pt(9.5)
            r_con.font.color.rgb = RGBColor(100, 116, 139)

    lang = (settings.get("language") or "en").lower()
    section_titles = settings.get("section_titles") or {}

    for sec_key in order:
        if not visibility.get(sec_key, True):
            continue

        resolved_heading = resolve_section_title(sec_key, section_titles.get(sec_key), lang=lang).upper()

        if sec_key == "summary" and profile.get("summary"):
            sh = doc.add_heading(resolved_heading, level=1)
            sh.runs[0].font.color.rgb = theme_rgb
            doc.add_paragraph(profile["summary"])

        elif sec_key == "skills" and structured_data.get("skills"):
            sh = doc.add_heading(resolved_heading, level=1)
            sh.runs[0].font.color.rgb = theme_rgb
            skill_items = [
                format_skill_display(s.get("name", ""), s.get("proficiency"), lang=lang)
                for s in structured_data["skills"]
                if s.get("name")
            ]
            skill_items = [si for si in skill_items if si]
            doc.add_paragraph(", ".join(skill_items))

        elif sec_key == "experience" and structured_data.get("experience"):
            sh = doc.add_heading(resolved_heading, level=1)
            sh.runs[0].font.color.rgb = theme_rgb
            for exp in structured_data["experience"]:
                title = exp.get("title", "")
                company = exp.get("company", "")
                p_exp = doc.add_paragraph()
                r_title = p_exp.add_run(f"{title} | {company}")
                r_title.bold = True
                date_range = f"{exp.get('start_date', '')} – {exp.get('end_date', '')}".strip(" –")
                if date_range:
                    p_exp.add_run(f"\n{date_range}").italic = True
                if exp.get("description"):
                    for line in exp["description"].splitlines():
                        if line.strip():
                            doc.add_paragraph(line.strip().lstrip("-•* "), style="List Bullet")
                for hl in exp.get("highlights", []):
                    if hl and hl.strip():
                        doc.add_paragraph(hl.strip(), style="List Bullet")

        elif sec_key == "education" and structured_data.get("education"):
            sh = doc.add_heading(resolved_heading, level=1)
            sh.runs[0].font.color.rgb = theme_rgb
            for edu in structured_data["education"]:
                p_edu = doc.add_paragraph()
                p_edu.add_run(f"{edu.get('degree', '')} in {edu.get('field', '')} — {edu.get('institution', '')}").bold = True
                date_range = f"{edu.get('start_date', '')} – {edu.get('end_date', '')}".strip(" –")
                if date_range:
                    p_edu.add_run(f"\n{date_range}").italic = True
                if edu.get("description"):
                    doc.add_paragraph(edu["description"])

        elif sec_key == "projects" and structured_data.get("projects"):
            sh = doc.add_heading(resolved_heading, level=1)
            sh.runs[0].font.color.rgb = theme_rgb
            for proj in structured_data["projects"]:
                p_proj = doc.add_paragraph()
                p_proj.add_run(proj.get("name", "")).bold = True
                if proj.get("technologies"):
                    p_proj.add_run(f" ({proj['technologies']})").italic = True
                if proj.get("description"):
                    doc.add_paragraph(proj["description"])

        elif sec_key == "certifications" and structured_data.get("certifications"):
            sh = doc.add_heading(resolved_heading, level=1)
            sh.runs[0].font.color.rgb = theme_rgb
            for cert in structured_data["certifications"]:
                text = f"{cert.get('name', '')} — {cert.get('issuer', '')} ({cert.get('issue_date', '')})"
                doc.add_paragraph(text, style="List Bullet")

        elif sec_key == "languages" and structured_data.get("languages"):
            sh = doc.add_heading(resolved_heading, level=1)
            sh.runs[0].font.color.rgb = theme_rgb
            items = [
                format_language_display(l.get("language", ""), l.get("proficiency"), lang=lang)
                for l in structured_data["languages"]
                if l.get("language")
            ]
            items = [it for it in items if it]
            doc.add_paragraph(", ".join(items))


    doc.save(str(dest_path))


# ---------------------------------------------------------------------------
# Document export helpers (motivation letter / cover letter)
# Reuse existing ReportLab + python-docx infrastructure. No new dependencies.
# ---------------------------------------------------------------------------

def render_document_pdf(content: str, title: str, dest_path: Path) -> None:
    """
    Render plain text content as a cleanly formatted PDF document.
    Used for motivation letters, cold emails, and other generated materials.
    """
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontSize=16,
        spaceAfter=16,
        textColor=colors.HexColor("#1e3a5f"),
        fontName="Helvetica-Bold",
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontSize=11,
        leading=17,
        spaceAfter=10,
        textColor=colors.HexColor("#1a1a2e"),
        fontName="Helvetica",
    )

    doc = SimpleDocTemplate(
        str(dest_path),
        pagesize=A4,
        topMargin=54,
        bottomMargin=54,
        leftMargin=60,
        rightMargin=60,
    )

    story = []
    if title:
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 8))

    for para in content.split("\n\n"):
        para = para.strip()
        if not para:
            story.append(Spacer(1, 6))
            continue
        # Escape XML special chars for ReportLab
        safe = para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        # Preserve single newlines as <br/>
        safe = safe.replace("\n", "<br/>")
        story.append(Paragraph(safe, body_style))
        story.append(Spacer(1, 4))

    doc.build(story)
    validate_pdf_file(dest_path)


def render_document_docx(content: str, title: str, dest_path: Path) -> None:
    """
    Render plain text content as a cleanly formatted DOCX document.
    Used for motivation letters, cold emails, and other generated materials.
    """
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    if title:
        h = doc.add_heading(title, level=1)
        h.runs[0].font.color.rgb = RGBColor(0x1e, 0x3a, 0x5f)
        h.runs[0].font.size = Pt(16)
        doc.add_paragraph("")

    # Body paragraphs
    for para in content.split("\n\n"):
        para = para.strip()
        if not para:
            doc.add_paragraph("")
            continue
        p = doc.add_paragraph()
        for i, line in enumerate(para.split("\n")):
            if i > 0:
                p.add_run("\n")
            run = p.add_run(line)
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

    doc.save(str(dest_path))
