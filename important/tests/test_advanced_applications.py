import io
from datetime import date, datetime, timedelta
from fastapi.testclient import TestClient
import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from src.app.database import get_db
from src.app.main import app
from src.app.models import Base, Company, Country, Job


def test_advanced_applications_contacts_and_timeline() -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)

    with session_factory() as session:
        country = Country(id=1, name="Morocco", code="MA")
        company = Company(id=1, name="CyberSec Inc", website_url="https://cybersec.test", active=True)
        session.add_all([country, company])
        session.flush()

        job = Job(
            id=101,
            company_id=1,
            country_id=1,
            title="SOC Analyst Intern",
            job_url="https://cybersec.test/jobs/101",
            active=True,
            description="We are seeking a SOC analyst with Python, Linux, and SIEM skills.",
            discovered_at=datetime.now(),
            created_at=datetime.now(),
            updated_at=datetime.now(),
        )
        session.add(job)
        session.commit()

    def fake_db():
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = fake_db
    client = TestClient(app)

    # 1. Register candidate
    reg_resp = client.post(
        "/api/v1/auth/register",
        json={"email": "sarah.candidate@test.com", "password": "Password123!"},
    )
    assert reg_resp.status_code == 201
    token = reg_resp.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 2. Create a contact directly
    contact_resp = client.post(
        "/api/v1/contacts",
        headers=headers,
        json={
            "name": "Sarah Johnson",
            "email": "sarah.johnson@cybersec.test",
            "company": "CyberSec Inc",
            "role": "Lead Technical Recruiter",
            "phone": "+1-555-0199",
            "linkedin_url": "https://linkedin.com/in/sarah-johnson-recruiter",
            "tags": "recruiter, tech",
            "notes": "Met at cybersecurity career fair",
        },
    )
    assert contact_resp.status_code == 201
    contact = contact_resp.json()
    assert contact["name"] == "Sarah Johnson"
    assert contact["email"] == "sarah.johnson@cybersec.test"
    contact_id = contact["id"]

    # 3. Create a manual application with smart contact linking
    # Passing the exact contact email should auto-link contact_id!
    app_resp = client.post(
        "/api/v1/applications",
        headers=headers,
        json={
            "job_title": "Security Analyst",
            "company_name": "CyberSec Inc",
            "source": "LinkedIn",
            "status": "applied",
            "contact_email": "SARAH.JOHNSON@cybersec.test",  # Uppercase to test normalization
            "next_follow_up": (date.today() + timedelta(days=5)).isoformat(),
        },
    )
    assert app_resp.status_code == 201
    app_data = app_resp.json()
    app_id = app_data["id"]
    assert app_data["contact_id"] == contact_id

    # 4. Check contact details now shows 1 related application
    contact_get = client.get(f"/api/v1/contacts/{contact_id}", headers=headers)
    assert contact_get.status_code == 200
    c_detail = contact_get.json()
    assert c_detail["applications_count"] == 1
    assert len(c_detail["related_applications"]) == 1
    assert c_detail["related_applications"][0]["job_title"] == "Security Analyst"

    # 5. Log contact interaction
    inter_resp = client.post(
        f"/api/v1/contacts/{contact_id}/interactions",
        headers=headers,
        json={
            "type": "email_sent",
            "summary": "Followed up regarding Security Analyst opening",
            "details": "Sent revised resume and project link",
            "application_id": app_id,
        },
    )
    assert inter_resp.status_code == 201
    assert inter_resp.json()["summary"] == "Followed up regarding Security Analyst opening"

    # 6. Check duplicate contact detection endpoint
    dup_resp = client.get(
        f"/api/v1/contacts/check-duplicates?email=sarah.johnson@cybersec.test",
        headers=headers,
    )
    assert dup_resp.status_code == 200
    dups = dup_resp.json()
    assert len(dups) == 1
    assert dups[0]["confidence"] == "exact"

    # 7. Add custom timeline event on application
    evt_resp = client.post(
        f"/api/v1/applications/{app_id}/events",
        headers=headers,
        json={
            "event_type": "note",
            "event_title": "Completed technical challenge",
            "notes": "Submitted Python script with 100% test coverage",
        },
    )
    assert evt_resp.status_code == 201
    evt = evt_resp.json()
    assert evt["event_title"] == "Completed technical challenge"

    # 8. Follow-up quick actions (snooze, complete)
    fu_snooze = client.post(
        f"/api/v1/applications/{app_id}/follow-up",
        headers=headers,
        json={"action": "snooze", "days": 3},
    )
    assert fu_snooze.status_code == 200
    expected_snooze_date = (date.today() + timedelta(days=8)).isoformat()
    assert fu_snooze.json()["next_follow_up"] == expected_snooze_date

    fu_complete = client.post(
        f"/api/v1/applications/{app_id}/follow-up",
        headers=headers,
        json={"action": "complete", "notes": "Got reply from Sarah"},
    )
    assert fu_complete.status_code == 200
    assert fu_complete.json()["next_follow_up"] is None

    # 9. Test application intelligence endpoint
    intel_resp = client.get(f"/api/v1/applications/{app_id}/intelligence", headers=headers)
    assert intel_resp.status_code == 200
    intel = intel_resp.json()
    assert "match_score" in intel
    assert "verdict" in intel
    assert "strong_matches" in intel
    assert "potential_gaps" in intel

    # 10. Test application stats with new rich analytics
    stats_resp = client.get("/api/v1/applications/stats", headers=headers)
    assert stats_resp.status_code == 200
    stats = stats_resp.json()
    assert "source_analytics" in stats
    assert len(stats["source_analytics"]) > 0
    assert stats["source_analytics"][0]["source"] == "LinkedIn"
    assert "activity_timeline" in stats
    assert "7d" in stats["activity_timeline"]
    assert "overview" in stats
    assert stats["overview"]["total_applications"] >= 1

    # 11. Test contacts export XLSX
    export_resp = client.get("/api/v1/contacts/export/xlsx", headers=headers)
    assert export_resp.status_code == 200
    assert export_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    xlsx_bytes = export_resp.content
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws.title == "Contacts"
    assert ws.cell(row=1, column=3).value == "Full Name"
    assert ws.cell(row=2, column=3).value == "Sarah Johnson"

    # 12. Test contacts import preview
    import_file = io.BytesIO()
    import_wb = openpyxl.Workbook()
    import_ws = import_wb.active
    import_ws.append(["Name", "Email Address", "Company Name", "Role", "Phone"])
    import_ws.append(["Alex Smith", "alex@techcorp.test", "TechCorp", "Talent Partner", "+1-555-8888"])
    import_ws.append(["Sarah Johnson", "sarah.johnson@cybersec.test", "CyberSec Inc", "Lead Recruiter", ""])
    import_wb.save(import_file)
    import_file.seek(0)

    preview_resp = client.post(
        "/api/v1/contacts/preview-import",
        headers=headers,
        files={"file": ("candidates.xlsx", import_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert preview_resp.status_code == 200
    preview = preview_resp.json()
    assert preview["total_detected"] == 2
    assert preview["new_count"] == 1
    assert preview["existing_count"] == 1

    # 13. Confirm import with skip duplicates
    confirm_resp = client.post(
        "/api/v1/contacts/confirm-import",
        headers=headers,
        json={
            "mode": "skip_duplicates",
            "rows": preview["preview_rows"],
        },
    )
    assert confirm_resp.status_code == 200
    res = confirm_resp.json()
    assert res["created_count"] == 1
    assert res["skipped_count"] == 1

    # Verify newly imported contact exists
    contacts_list = client.get("/api/v1/contacts", headers=headers).json()
    assert len(contacts_list) == 2
    names = {c["name"] for c in contacts_list}
    assert "Alex Smith" in names
