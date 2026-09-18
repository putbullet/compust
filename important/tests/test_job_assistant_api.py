"""
test_job_assistant_api.py
=========================
API integration tests for the job-target endpoints.
Ollama is mocked - no network calls.
"""
from datetime import datetime, timezone
from unittest.mock import patch, MagicMock
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from src.app.database import get_db
from src.app.main import app
from src.app.models import Base, Company, Country, Job, User, Resume
from src.app.security import create_access_token


# ---------------------------------------------------------------------------
# Shared test environment fixture
# ---------------------------------------------------------------------------

@pytest.fixture
def api_env():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)

    with factory() as session:
        now = datetime.now(timezone.utc).replace(tzinfo=None)
        user = User(
            id=1, email="ana@compust.ai", password_hash="pw",
            is_active=True, first_name="Ana", last_name="Silva",
            created_at=now, updated_at=now,
        )
        session.add(user)

        other_user = User(
            id=2, email="other@compust.ai", password_hash="pw",
            is_active=True, first_name="Other", last_name="User",
            created_at=now, updated_at=now,
        )
        session.add(other_user)

        master_resume = Resume(
            id=10, user_id=1,
            title="Master Resume",
            is_active=True, is_default=True,
            structured_data={
                "profile": {
                    "full_name": "Ana Silva",
                    "headline": "Software Engineer",
                    "email": "ana@compust.ai",
                    "summary": "Experienced Python developer with FastAPI and Docker knowledge.",
                },
                "skills": [
                    {"id": "sk1", "name": "Python", "category": "Backend", "proficiency": "Expert"},
                    {"id": "sk2", "name": "Docker", "category": "DevOps", "proficiency": "Intermediate"},
                    {"id": "sk3", "name": "FastAPI", "category": "Backend", "proficiency": "Advanced"},
                ],
                "experience": [
                    {
                        "id": "exp1",
                        "company": "StartupXYZ",
                        "title": "Software Engineer",
                        "description": "Built Python REST APIs using FastAPI and deployed with Docker.",
                        "highlights": ["Improved API response time by 40%"],
                        "start_date": "2023-01", "end_date": "Present",
                    }
                ],
                "education": [
                    {"id": "edu1", "institution": "ENSIAS", "degree": "BSc", "field": "Software Engineering"},
                ],
                "projects": [], "languages": [{"id": "l1", "language": "English", "proficiency": "Fluent"}],
                "certifications": [], "custom_sections": [],
            },
            settings={"template": "modern", "theme_color": "#2563eb"},
            created_at=now, updated_at=now,
        )
        session.add(master_resume)
        session.commit()

    def override_get_db():
        with factory() as s:
            yield s

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)
    token = create_access_token({"sub": "1"})
    token_other = create_access_token({"sub": "2"})
    yield client, token, token_other
    app.dependency_overrides.clear()


VALID_JOB_PAYLOAD = {
    "target_role": "Senior Python Engineer",
    "job_description": (
        "We are looking for a Senior Python Engineer with extensive Docker and FastAPI experience. "
        "The role requires strong knowledge of REST API design, microservices architecture, "
        "and containerized deployments. PostgreSQL and Kubernetes are required. " * 5
    ),
    "language": "en",
}

OLLAMA_UNAVAILABLE = {"status": "disconnected", "models": [], "selected_model": None}


# ---------------------------------------------------------------------------
# Analyze endpoint — success cases
# ---------------------------------------------------------------------------

def test_analyze_endpoint_valid_returns_200(api_env):
    client, token, _ = api_env
    with patch("src.app.services.job_assistant.check_ollama_runtime", return_value=OLLAMA_UNAVAILABLE):
        resp = client.post(
            "/api/v1/resumes/10/job-target/analyze",
            json=VALID_JOB_PAYLOAD,
            headers={"Authorization": f"Bearer {token}"},
        )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["resume_id"] == 10
    assert data["target_role"] == "Senior Python Engineer"
    assert data["language"] == "en"
    assert isinstance(data["job_requirements"], list)
    assert isinstance(data["resume_recommendations"], list)
    assert "deterministic_score" in data


def test_analyze_endpoint_returns_deterministic_score(api_env):
    client, token, _ = api_env
    with patch("src.app.services.job_assistant.check_ollama_runtime", return_value=OLLAMA_UNAVAILABLE):
        resp = client.post(
            "/api/v1/resumes/10/job-target/analyze",
            json=VALID_JOB_PAYLOAD,
            headers={"Authorization": f"Bearer {token}"},
        )
    data = resp.json()
    score = data["deterministic_score"]
    assert 0.0 <= score["overall"] <= 10.0
    assert 0.0 <= score["technical_skills"] <= 10.0


def test_analyze_endpoint_flags_missing_skills(api_env):
    """Kubernetes is not in candidate data - must appear in missing_or_unconfirmed."""
    client, token, _ = api_env
    with patch("src.app.services.job_assistant.check_ollama_runtime", return_value=OLLAMA_UNAVAILABLE):
        resp = client.post(
            "/api/v1/resumes/10/job-target/analyze",
            json=VALID_JOB_PAYLOAD,
            headers={"Authorization": f"Bearer {token}"},
        )
    data = resp.json()
    missing = [m["requirement"].lower() for m in data["match_analysis"]["missing_or_unconfirmed"]]
    assert any("kubernetes" in r for r in missing)


def test_analyze_endpoint_finds_strong_matches(api_env):
    """Python and Docker are in candidate data - must appear in strong matches."""
    client, token, _ = api_env
    with patch("src.app.services.job_assistant.check_ollama_runtime", return_value=OLLAMA_UNAVAILABLE):
        resp = client.post(
            "/api/v1/resumes/10/job-target/analyze",
            json=VALID_JOB_PAYLOAD,
            headers={"Authorization": f"Bearer {token}"},
        )
    data = resp.json()
    strong = [m["requirement"].lower() for m in data["match_analysis"]["strong_matches"]]
    assert any("python" in r for r in strong)
    assert any("docker" in r for r in strong)


# ---------------------------------------------------------------------------
# Analyze endpoint — validation failures
# ---------------------------------------------------------------------------

def test_analyze_endpoint_missing_role_returns_422(api_env):
    client, token, _ = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/analyze",
        json={"target_role": "", "job_description": "x" * 100, "language": "en"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 422


def test_analyze_endpoint_missing_description_returns_422(api_env):
    client, token, _ = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/analyze",
        json={"target_role": "Engineer", "job_description": "too short", "language": "en"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 422


def test_analyze_endpoint_oversized_description_returns_422(api_env):
    client, token, _ = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/analyze",
        json={"target_role": "Engineer", "job_description": "x" * 12001, "language": "en"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 422


# ---------------------------------------------------------------------------
# Authorization
# ---------------------------------------------------------------------------

def test_analyze_endpoint_unauthenticated_returns_401(api_env):
    client, _, _ = api_env
    resp = client.post("/api/v1/resumes/10/job-target/analyze", json=VALID_JOB_PAYLOAD)
    assert resp.status_code == 401


def test_analyze_wrong_users_resume_returns_404(api_env):
    """User 2 must not access user 1's resume."""
    client, _, token_other = api_env
    with patch("src.app.services.job_assistant.check_ollama_runtime", return_value=OLLAMA_UNAVAILABLE):
        resp = client.post(
            "/api/v1/resumes/10/job-target/analyze",
            json=VALID_JOB_PAYLOAD,
            headers={"Authorization": f"Bearer {token_other}"},
        )
    assert resp.status_code == 404


# ---------------------------------------------------------------------------
# Save tailored copy - non-destructive safety
# ---------------------------------------------------------------------------

def test_save_tailored_copy_creates_new_resume(api_env):
    client, token, _ = api_env
    headers = {"Authorization": f"Bearer {token}"}
    resp = client.post(
        "/api/v1/resumes/10/job-target/save-tailored-copy",
        json={
            "new_title": "Python Engineer Tailored Copy",
            "accepted_recommendation_ids": [],
            "accepted_recommendations": [],
        },
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["id"] != 10
    assert data["title"] == "Python Engineer Tailored Copy"
    assert data["source_resume_id"] == 10
    assert data["is_default"] is False


def test_save_tailored_copy_master_resume_unchanged(api_env):
    """CRITICAL: The master resume must be completely unchanged after creating a tailored copy."""
    client, token, _ = api_env
    headers = {"Authorization": f"Bearer {token}"}

    # Record master state before
    master_before = client.get("/api/v1/resumes/10", headers=headers).json()
    orig_summary = master_before["structured_data"]["profile"]["summary"]
    orig_title = master_before["title"]

    # Apply a summary rewrite recommendation
    resp = client.post(
        "/api/v1/resumes/10/job-target/save-tailored-copy",
        json={
            "new_title": "Tailored Copy",
            "accepted_recommendation_ids": ["r1"],
            "accepted_recommendations": [
                {
                    "id": "r1",
                    "rec_type": "rewrite",
                    "section": "summary",
                    "suggested_text": "Experienced Python engineer specializing in cloud infrastructure.",
                    "reason": "Better alignment with job",
                    "grounded": True,
                    "item_id": "",
                    "item_label": "",
                    "current_text": "",
                    "evidence_sources": [],
                }
            ],
        },
        headers=headers,
    )
    assert resp.status_code == 200, resp.text

    # Verify master is unchanged
    master_after = client.get("/api/v1/resumes/10", headers=headers).json()
    assert master_after["title"] == orig_title
    assert master_after["structured_data"]["profile"]["summary"] == orig_summary
    assert master_after["is_default"] is True
    assert master_after["source_resume_id"] is None


def test_save_tailored_copy_has_source_resume_id(api_env):
    client, token, _ = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/save-tailored-copy",
        json={"new_title": "Copy", "accepted_recommendation_ids": [], "accepted_recommendations": []},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.json()["source_resume_id"] == 10


def test_save_tailored_copy_unauthorized_returns_401(api_env):
    client, _, _ = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/save-tailored-copy",
        json={"new_title": "Copy", "accepted_recommendation_ids": [], "accepted_recommendations": []},
    )
    assert resp.status_code == 401


def test_save_tailored_copy_wrong_user_returns_404(api_env):
    client, _, token_other = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/save-tailored-copy",
        json={"new_title": "Copy", "accepted_recommendation_ids": [], "accepted_recommendations": []},
        headers={"Authorization": f"Bearer {token_other}"},
    )
    assert resp.status_code == 404


# ---------------------------------------------------------------------------
# Export letter endpoint
# ---------------------------------------------------------------------------

def test_export_letter_unauthorized_returns_401(api_env):
    client, _, _ = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/export-letter",
        json={"content": "Dear Hiring Manager,\n\nI am writing...\n\nBest regards,\nAna.", "file_format": "pdf", "suggested_filename": "motivation_letter"},
    )
    assert resp.status_code == 401


def test_export_letter_wrong_user_returns_404(api_env):
    client, _, token_other = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/export-letter",
        json={"content": "Dear Hiring Manager,\n\nI am writing...\n\nBest regards,\nAna.", "file_format": "pdf", "suggested_filename": "motivation_letter"},
        headers={"Authorization": f"Bearer {token_other}"},
    )
    assert resp.status_code == 404


def test_export_letter_success_pdf(api_env):
    import io
    from pypdf import PdfReader
    client, token, _ = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/export-letter",
        json={
            "content": "Dear Hiring Manager,\n\nI am writing to express my strong enthusiasm for the role.\n\nSincerely,\nAna Silva",
            "file_format": "pdf",
            "suggested_filename": "Ana_Silva_Motivation_Letter",
        },
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert "attachment" in resp.headers["content-disposition"]
    assert "Ana_Silva_Motivation_Letter" in resp.headers["content-disposition"]
    assert resp.content.startswith(b"%PDF-")
    reader = PdfReader(io.BytesIO(resp.content))
    assert len(reader.pages) >= 1
    text = reader.pages[0].extract_text()
    assert "Ana Silva" in text or "Hiring Manager" in text


def test_export_letter_success_docx(api_env):
    import io
    from docx import Document
    client, token, _ = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/export-letter",
        json={
            "content": "Subject: Python Engineer Application\n\nDear Recruiter,\n\nI look forward to discussing the role.\n\nBest,\nAna",
            "file_format": "docx",
            "suggested_filename": "Cold_Email_Application",
        },
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 200
    assert "wordprocessingml.document" in resp.headers["content-type"]
    assert "attachment" in resp.headers["content-disposition"]
    assert resp.content.startswith(b"PK\x03\x04")
    doc = Document(io.BytesIO(resp.content))
    paragraphs_text = " ".join(p.text for p in doc.paragraphs)
    assert "Python Engineer Application" in paragraphs_text or "Dear Recruiter" in paragraphs_text


# ---------------------------------------------------------------------------
# Single material regeneration tests
# ---------------------------------------------------------------------------

def test_regenerate_material_unauthorized(api_env):
    client, _, _ = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/regenerate-material",
        json={
            "material_type": "cold_email",
            "target_role": "Senior Python Engineer",
            "job_description": "We need a senior python engineer with FastAPI and Docker." * 3,
        },
    )
    assert resp.status_code == 401


def test_regenerate_material_wrong_owner(api_env):
    client, _, token_other = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/regenerate-material",
        json={
            "material_type": "cold_email",
            "target_role": "Senior Python Engineer",
            "job_description": "We need a senior python engineer with FastAPI and Docker." * 3,
        },
        headers={"Authorization": f"Bearer {token_other}"},
    )
    assert resp.status_code == 404


def test_regenerate_material_invalid_body(api_env):
    client, token, _ = api_env
    resp = client.post(
        "/api/v1/resumes/10/job-target/regenerate-material",
        json={
            "material_type": "cold_email",
            "target_role": "Senior Python Engineer",
            "job_description": "Short",  # under 50 chars
        },
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 422


def test_regenerate_material_ollama_offline(api_env):
    client, token, _ = api_env
    with patch("src.app.services.job_assistant.check_ollama_runtime") as mock_runtime:
        mock_runtime.return_value = {"status": "offline", "models": []}
        resp = client.post(
            "/api/v1/resumes/10/job-target/regenerate-material",
            json={
                "material_type": "cold_email",
                "target_role": "Senior Python Engineer",
                "job_description": "We need a senior python engineer with FastAPI and Docker experience." * 3,
            },
            headers={"Authorization": f"Bearer {token}"},
        )
    assert resp.status_code == 200
    data = resp.json()
    assert data["material_type"] == "cold_email"
    assert data["generated"] is False
    assert "Ollama" in data["error"]


def test_regenerate_material_success_mocked(api_env):
    import json
    client, token, _ = api_env
    mock_llm_response = json.dumps({
        "cold_email": {
            "subject": "Application for Senior Python Engineer - Ana Silva",
            "body": "Dear [RECRUITER NAME],\n\nI am writing to express my interest in the Senior Python Engineer role at [COMPANY NAME]. With strong experience in FastAPI and Docker, I am excited about this opportunity.\n\nSincerely,\nAna Silva",
            "placeholders": ["[RECRUITER NAME]", "[COMPANY NAME]"],
        }
    })
    with patch("src.app.services.job_assistant.check_ollama_runtime") as mock_runtime, \
         patch("src.app.services.job_assistant.generate_completion") as mock_generate:
        mock_runtime.return_value = {"status": "connected", "models": ["qwen3.5:0.8b"], "selected_model": "qwen3.5:0.8b"}
        mock_generate.return_value = mock_llm_response

        resp = client.post(
            "/api/v1/resumes/10/job-target/regenerate-material",
            json={
                "material_type": "cold_email",
                "target_role": "Senior Python Engineer",
                "job_description": "We need a senior python engineer with FastAPI and Docker experience." * 3,
            },
            headers={"Authorization": f"Bearer {token}"},
        )
    assert resp.status_code == 200
    data = resp.json()
    assert data["material_type"] == "cold_email"
    assert data["generated"] is True
    assert "Ana Silva" in data["subject"]
    assert "[COMPANY NAME]" in data["body"]
    assert "[RECRUITER NAME]" in data["placeholders"]

