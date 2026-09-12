import io
from datetime import datetime, timezone
import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from src.app.database import get_db
from src.app.main import app
from src.app.models import Base, Company, Country, Job


def setup_test_client():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)

    now = datetime.now(timezone.utc).replace(tzinfo=None)
    with session_factory() as session:
        country = Country(id=1, name="France", code="FR")
        company = Company(id=1, name="Stripe France", website_url="https://stripe.com", active=True)
        session.add_all([country, company])
        session.flush()

        job = Job(
            id=101,
            company_id=1,
            country_id=1,
            title="Senior Python Architect",
            job_url="https://stripe.com/jobs/101",
            location="Paris, France",
            active=True,
            discovered_at=now,
            created_at=now,
            updated_at=now,
        )
        session.add(job)
        session.commit()

    def fake_db():
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = fake_db
    return TestClient(app)


def test_unified_manual_and_compust_applications():
    client = setup_test_client()

    # 1. Register candidate user
    reg = client.post(
        "/api/v1/auth/register",
        json={"email": "tracker_user@test.com", "password": "Password123!"},
    )
    assert reg.status_code == 201
    token = reg.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 2. Add manual application (LinkedIn)
    manual_payload = {
        "job_title": "Staff Engineer",
        "company_name": "Datadog",
        "source": "LinkedIn",
        "status": "applied",
        "location": "Remote",
        "country": "France",
        "salary": "€105k",
        "priority": "high",
        "contact_name": "Alice Recruiter",
        "contact_email": "alice@datadog.com",
        "notes": "Applied with revised resume targeting observability.",
    }
    res = client.post("/api/v1/applications", json=manual_payload, headers=headers)
    assert res.status_code == 201
    data = res.json()
    assert data["job_id"] is None
    assert data["effective_title"] == "Staff Engineer"
    assert data["effective_company"] == "Datadog"
    assert data["source"] == "LinkedIn"
    assert data["status"] == "applied"
    manual_app_id = data["id"]
    assert len(data["history"]) == 1
    assert data["history"][0]["to_status"] == "applied"

    # 3. Add automatic Compust application from Opportunities
    compust_payload = {"status": "saved", "notes": "Found via Compust Opportunities"}
    res_comp = client.post("/api/v1/applications/101", json=compust_payload, headers=headers)
    assert res_comp.status_code == 201
    data_comp = res_comp.json()
    assert data_comp["job_id"] == 101
    assert data_comp["effective_title"] == "Senior Python Architect"
    assert data_comp["effective_company"] == "Stripe France"
    assert data_comp["source"] == "Compust"
    assert data_comp["status"] == "saved"
    compust_app_id = data_comp["id"]

    # 4. Progress stages and verify history logging
    # Stage transition 1: applied -> 1st_interview
    patch_1 = client.patch(
        f"/api/v1/applications/{manual_app_id}",
        json={"status": "1st_interview", "notes": "HR screening scheduled"},
        headers=headers,
    )
    assert patch_1.status_code == 200
    assert patch_1.json()["status"] == "1st_interview"
    assert len(patch_1.json()["history"]) == 2
    assert patch_1.json()["history"][-1]["from_status"] == "applied"
    assert patch_1.json()["history"][-1]["to_status"] == "1st_interview"

    # Stage transition 2: 1st_interview -> 2nd_interview
    patch_2 = client.patch(
        f"/api/v1/applications/{manual_app_id}",
        json={"status": "2nd_interview"},
        headers=headers,
    )
    assert patch_2.status_code == 200
    assert len(patch_2.json()["history"]) == 3
    assert patch_2.json()["history"][-1]["from_status"] == "1st_interview"
    assert patch_2.json()["history"][-1]["to_status"] == "2nd_interview"

    # Stage transition 3: 2nd_interview -> offer
    patch_3 = client.patch(
        f"/api/v1/applications/{manual_app_id}",
        json={"status": "offer", "notes": "Offer letter received!"},
        headers=headers,
    )
    assert patch_3.status_code == 200
    assert patch_3.json()["status"] == "offer"

    # Stage transition 4: offer -> accepted
    patch_4 = client.patch(
        f"/api/v1/applications/{manual_app_id}",
        json={"status": "accepted"},
        headers=headers,
    )
    assert patch_4.status_code == 200
    assert patch_4.json()["status"] == "accepted"

    # 5. Test Filters and Search
    # Search by company name
    res_search = client.get("/api/v1/applications?search=Datadog", headers=headers)
    assert res_search.status_code == 200
    assert len(res_search.json()) == 1
    assert res_search.json()[0]["id"] == manual_app_id

    # Filter by source
    res_src = client.get("/api/v1/applications?source=Compust", headers=headers)
    assert res_src.status_code == 200
    assert len(res_src.json()) == 1
    assert res_src.json()[0]["id"] == compust_app_id

    # 6. Test Reliable Statistics Calculation
    stats_res = client.get("/api/v1/applications/stats", headers=headers)
    assert stats_res.status_code == 200
    stats = stats_res.json()
    assert stats["total_applications"] == 2
    assert stats["source_breakdown"]["LinkedIn"] == 1
    assert stats["source_breakdown"]["Compust"] == 1
    assert stats["interview_rate_percent"] > 0
    assert stats["offer_rate_percent"] > 0

    # 7. Test Dynamic Sankey Pipeline Funnel Endpoint
    pipeline_res = client.get("/api/v1/applications/pipeline", headers=headers)
    assert pipeline_res.status_code == 200
    pipeline = pipeline_res.json()
    assert pipeline["total"] == 2
    assert len(pipeline["nodes"]) > 0
    assert len(pipeline["links"]) > 0
    node_ids = {n["id"] for n in pipeline["nodes"]}
    assert "stage_applications" in node_ids

    # 8. Test Excel Workbook Export (.xlsx)
    export_res = client.get("/api/v1/applications/export", headers=headers)
    assert export_res.status_code == 200
    assert (
        export_res.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    # Parse returned workbook bytes with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(export_res.content))
    assert "Applications" in wb.sheetnames
    ws = wb["Applications"]
    assert ws.max_row >= 3  # Header + 2 data rows
    headers_row = [cell.value for cell in ws[1]]
    assert "Job Title" in headers_row
    assert "Company" in headers_row
    assert "Source" in headers_row
    assert "Status" in headers_row

    # 9. Test Security and Isolation
    reg2 = client.post(
        "/api/v1/auth/register",
        json={"email": "attacker@test.com", "password": "Password123!"},
    )
    attacker_token = reg2.json()["access_token"]
    attacker_headers = {"Authorization": f"Bearer {attacker_token}"}

    # Attacker cannot modify user's application
    patch_attack = client.patch(
        f"/api/v1/applications/{manual_app_id}",
        json={"status": "rejected"},
        headers=attacker_headers,
    )
    assert patch_attack.status_code == 404

    # Attacker cannot delete user's application
    del_attack = client.delete(
        f"/api/v1/applications/{manual_app_id}",
        headers=attacker_headers,
    )
    assert del_attack.status_code == 404

    # Attacker sees empty applications
    assert client.get("/api/v1/applications", headers=attacker_headers).json() == []
